"""Builds Amma ji's recipe sheet: the INDB dishes that are also in Mendeley, with steps written for the app.

Joins three sources:
  app/data/indb_recipes.json      ingredients, amounts, nutrition, diet, allergens  (scripts/import_indb.py)
  app/data/mendeley_matches.json  prep/cook time, course, cuisine, credit link       (scripts/match_mendeley.py)
  scripts/steps/batch_*.json      steps in 4 languages, timers, safety, mistakes, swaps (written for Amma ji)

Writes:
  app/data/catalogue.json         what the app reads
  ../data/amma_ji_recipes.xlsx    the same data as a workbook, one sheet per table
  ../data/csv/*.csv               one CSV per table, ready to import into Supabase

Run from the backend folder:  .venv/bin/python scripts/build_catalogue.py
"""
import csv
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

BACKEND = Path(__file__).resolve().parent.parent
DATA = BACKEND / "app" / "data"
STEPS = BACKEND / "scripts" / "steps"
SHEET = BACKEND.parent / "data"
LANGS = ("en", "hi-en", "hi", "pa")
CREDIT = ("Ingredients and nutrition: Indian Nutrient Databank (INDB), Vijayakumar A, Dubasi HB, Awasthi A, "
          "Jaacks LM, Curr Dev Nutr 2024;8(7):103790. Nutrition values are estimates. Method adapted from "
          "Archana's Kitchen (archanaskitchen.com) via the 6000+ Indian Food Recipes Dataset, Kanishka Jain, "
          "Mendeley Data 2020, doi:10.17632/xsphgmmh7b.1, CC BY 4.0, rewritten in Amma ji's words.")


def load(path):
    return json.loads(path.read_text(encoding="utf-8"))


def uses(step_text, ingredients, catalogue):
    """Which ingredients a step mentions, so finishing the step can mark them as added."""
    text = step_text.lower()
    found = []
    for i in ingredients:
        known = catalogue.get(i["id"], {})
        words = {i["as_written"].lower(), known.get("names", {}).get("en", ""), *known.get("spellings", [])} - {""}
        if any(re.search(rf"\b{re.escape(w)}", text) for w in words):
            found.append(i["id"])
    return sorted(set(found))


def main():
    indb = {r["id"]: r for r in load(DATA / "indb_recipes.json")["recipes"]}
    matches = load(DATA / "mendeley_matches.json")["matches"]
    known = {i["id"]: i for i in load(DATA / "ingredients.json")}
    safety = load(STEPS / "safety.json")
    written = {}
    for path in sorted(STEPS.glob("batch_*.json")):
        written.update(load(path))

    problems, recipes = [], []
    for code in sorted(matches):
        dish, match = indb[code], matches[code]
        steps_src = written.get(code)
        swaps = (steps_src or {}).get("swaps", {})
        ingredients = []
        for i in dish["ingredients"]:
            if any(x["id"] == i["id"] and x["as_written"] == i["as_written"] for x in ingredients):
                continue
            ingredients.append({**{k: i[k] for k in ("id", "as_written", "qty", "unit", "essential")},
                                "substitutes": swaps.get(i["id"], [])})
        for old, new in swaps.items():
            if old not in {i["id"] for i in ingredients}:
                problems.append(f"{code} swap for {old}, which the dish does not use")
            for x in new:
                if x != "skip" and x not in known:
                    problems.append(f"{code} swap {old} -> {x}: unknown ingredient")
        steps = []
        for n, s in enumerate((steps_src or {}).get("steps", []), 1):
            missing = [lang for lang in LANGS if not s.get(lang)]
            if missing:
                problems.append(f"{code} step {n} has no {', '.join(missing)}")
            if s.get("safety") and s["safety"] not in safety:
                problems.append(f"{code} step {n}: unknown safety note {s['safety']}")
            steps.append({
                "n": n,
                "text": {lang: s.get(lang, "") for lang in LANGS},
                "uses": uses(s.get("en", ""), ingredients, known),
                "timer_min": s.get("timer"),
                "safety": safety.get(s.get("safety")) if s.get("safety") else None,
            })
        local = (steps_src or {}).get("names", {})
        recipes.append({
            "id": code,
            "name": {"en": dish["name"], "hi-en": local.get("hi-en") or dish["local_name"] or dish["name"],
                     "hi": local.get("hi") or dish["local_name"] or dish["name"],
                     "pa": local.get("pa") or dish["local_name"] or dish["name"]},
            "kind": dish["kind"],
            "diet": dish["diet"],
            "allergens": dish["allergens"],
            "servings": dish["servings"],
            "prep_min": match["prep_min"],
            "cook_min": match["cook_min"],
            "total_min": match["total_min"],
            "course": match["course"],
            "cuisine": match["cuisine"],
            "ingredients": ingredients,
            "steps": steps,
            "has_steps": bool(steps),
            "common_mistakes": [{"mistake": m["mistake"], "fix": {"en": m["en"], "hi": m.get("hi", "")}}
                                for m in (steps_src or {}).get("mistakes", [])],
            "nutrition_per_serving": dish["nutrition"],
            "source_note": CREDIT,
            "source_link": match["url"],
            "mendeley_name": match["mendeley_name"],
        })

    unknown = sorted(set(written) - set(matches))
    if unknown:
        problems.append(f"steps written for dishes not in the list: {', '.join(unknown)}")
    if problems:
        print("Problems:\n  " + "\n  ".join(problems))
        return 1

    (DATA / "catalogue.json").write_text(json.dumps({"credit": CREDIT, "recipes": recipes}, ensure_ascii=False,
                                                    separators=(",", ":")), encoding="utf-8")
    write_tables(recipes)
    ready = sum(r["has_steps"] for r in recipes)
    print(f"{len(recipes)} dishes in the sheet, {ready} with steps ({len(recipes) - ready} still to write)")
    return 0


def write_tables(recipes):
    """The same data as flat tables: an Excel workbook and one CSV per table for Supabase."""
    tables = {
        "dishes": [["id", "name_en", "name_hinglish", "name_hi", "name_pa", "kind", "diet", "allergens", "servings",
                    "prep_min", "cook_min", "total_min", "course", "cuisine", "kcal", "protein_g", "carbs_g", "fat_g",
                    "fibre_g", "nutrition_per", "has_steps", "source_link"]],
        "ingredients": [["dish_id", "ingredient_id", "as_written", "qty", "unit", "essential", "substitutes"]],
        "steps": [["dish_id", "n", "text_en", "text_hinglish", "text_hi", "text_pa", "timer_min", "safety_en"]],
        "mistakes": [["dish_id", "mistake", "fix_en", "fix_hi"]],
    }
    for r in recipes:
        n = r["nutrition_per_serving"]
        s = r["servings"]
        tables["dishes"].append([
            r["id"], r["name"]["en"], r["name"]["hi-en"], r["name"]["hi"], r["name"]["pa"], r["kind"], r["diet"],
            ", ".join(r["allergens"]), f"{s['count'] or ''} x {s['size'] or ''} {s['unit'] or ''}".strip(),
            r["prep_min"], r["cook_min"], r["total_min"], r["course"], r["cuisine"], n["kcal"], n["protein_g"],
            n["carbs_g"], n["fat_g"], n["fibre_g"], n["per"], "yes" if r["has_steps"] else "no", r["source_link"]])
        for i in r["ingredients"]:
            tables["ingredients"].append([r["id"], i["id"], i["as_written"], i["qty"], i["unit"],
                                          "yes" if i["essential"] else "no", ", ".join(i["substitutes"])])
        for st in r["steps"]:
            tables["steps"].append([r["id"], st["n"], st["text"]["en"], st["text"]["hi-en"], st["text"]["hi"],
                                    st["text"]["pa"], st["timer_min"], (st["safety"] or {}).get("en", "")])
        for m in r["common_mistakes"]:
            tables["mistakes"].append([r["id"], m["mistake"], m["fix"]["en"], m["fix"]["hi"]])

    (SHEET / "csv").mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.remove(book.active)
    for name, rows in tables.items():
        with open(SHEET / "csv" / f"{name}.csv", "w", newline="", encoding="utf-8") as f:
            csv.writer(f).writerows(rows)
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
    credits = book.create_sheet("credits")
    credits.append(["Amma ji recipe sheet"])
    credits.append([CREDIT])
    credits.append(["Steps, mistakes and dish names in Hindi and Punjabi were written for the app and are not "
                    "copied from either source."])
    book.save(SHEET / "amma_ji_recipes.xlsx")


if __name__ == "__main__":
    sys.exit(main())
